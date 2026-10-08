"""Pure document type detection, parsing, and deterministic text chunking."""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path


CHUNK_TARGET_CHARS = 1600
CHUNK_MAX_CHARS = 1800
CHUNK_OVERLAP_CHARS = 200


class DocumentParseError(ValueError):
    """Raised when a supported document cannot produce usable text."""


@dataclass
class ParsedDocument:
    text: str
    pages: list[tuple[int | None, str]]
    page_count: int | None


def detect_document_type(
    filename: str | None,
    mimetype: str | None,
    filetype: str | None,
) -> str | None:
    """Return pdf/docx/txt/md from Slack metadata, or None if unsupported."""
    mime = (mimetype or "").lower().split(";", 1)[0].strip()
    slack_type = (filetype or "").lower().strip()
    extension = Path(filename or "").suffix.lower()

    if mime == "application/pdf" or slack_type == "pdf" or extension == ".pdf":
        return "pdf"
    if (
        mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or slack_type == "docx"
        or extension == ".docx"
    ):
        return "docx"
    if mime == "text/markdown" or slack_type in {"markdown", "md"} or extension == ".md":
        return "md"
    if mime == "text/plain" or slack_type in {"text", "txt"} or extension == ".txt":
        return "txt"
    return None


def parse_document_bytes(data: bytes, document_type: str) -> ParsedDocument:
    """Extract full text and page-level text from a supported document."""
    if document_type == "pdf":
        return _parse_pdf(data)
    if document_type == "docx":
        return _parse_docx(data)
    if document_type in {"txt", "md"}:
        text = data.decode("utf-8", errors="replace")
        return ParsedDocument(text=text, pages=[(None, text)], page_count=None)
    raise DocumentParseError(f"Unsupported document type: {document_type}")


def _parse_pdf(data: bytes) -> ParsedDocument:
    import fitz

    try:
        with fitz.open(stream=data, filetype="pdf") as pdf:
            pages = [(index + 1, page.get_text("text")) for index, page in enumerate(pdf)]
            page_count = pdf.page_count
    except Exception as exc:
        raise DocumentParseError(f"PDF parsing failed: {exc}") from exc

    full_text = "\n\n".join(text for _, text in pages)
    if not full_text.strip():
        raise DocumentParseError(
            "PDF contains no extractable text; scanned/image-only PDFs require OCR, which is not supported."
        )
    return ParsedDocument(text=full_text, pages=pages, page_count=page_count)


def _parse_docx(data: bytes) -> ParsedDocument:
    from docx import Document as DocxDocument

    try:
        doc = DocxDocument(BytesIO(data))
        sections = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    sections.append(" | ".join(cells))
    except Exception as exc:
        raise DocumentParseError(f"DOCX parsing failed: {exc}") from exc

    text = "\n\n".join(sections)
    if not text.strip():
        raise DocumentParseError("DOCX contains no extractable paragraph or table text.")
    return ParsedDocument(text=text, pages=[(None, text)], page_count=None)


def chunk_parsed_document(parsed: ParsedDocument) -> list[dict]:
    """Split parsed text into deterministic, page-preserving chunks."""
    chunks = []
    for page_number, text in parsed.pages:
        for chunk_text in _chunk_text(text):
            chunks.append(
                {
                    "chunk_index": len(chunks),
                    "page_number": page_number,
                    "text": chunk_text,
                }
            )
    return chunks


def _chunk_text(text: str) -> list[str]:
    source = text.strip()
    if not source:
        return []

    chunks = []
    start = 0
    length = len(source)
    while start < length:
        remaining = length - start
        if remaining <= CHUNK_MAX_CHARS:
            end = length
        else:
            target = start + CHUNK_TARGET_CHARS
            minimum = start + 1200
            maximum = min(start + CHUNK_MAX_CHARS, length)
            before = _last_boundary(source, minimum, target)
            after = _first_boundary(source, target, maximum)
            end = before or after or maximum

        chunk = source[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= length:
            break
        next_start = max(0, end - CHUNK_OVERLAP_CHARS)
        start = next_start if next_start > start else end
    return chunks


def _last_boundary(text: str, minimum: int, target: int) -> int | None:
    for separator in ("\n\n", "\n", " "):
        index = text.rfind(separator, minimum, target + 1)
        if index >= minimum:
            return index + len(separator)
    return None


def _first_boundary(text: str, target: int, maximum: int) -> int | None:
    candidates = []
    for separator in ("\n\n", "\n", " "):
        index = text.find(separator, target, maximum)
        if index != -1:
            candidates.append(index + len(separator))
    return min(candidates) if candidates else None
