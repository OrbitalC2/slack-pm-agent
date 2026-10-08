import asyncio
import hashlib
import uuid
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import fitz
import pytest
from docx import Document as DocxDocument

from pm_agent.app.agents.chains import document_chunk_analysis_chain
from pm_agent.app.agents.schemas import (
    DocumentChunkAnalysisOutput,
    DocumentRequirement,
)
from pm_agent.app.db.models import Document, DocumentChunk
from pm_agent.app.services import document_service, gather_info_service
from pm_agent.app.utils.document_context import format_document_context
from pm_agent.app.utils.document_parsing import (
    CHUNK_MAX_CHARS,
    DocumentParseError,
    ParsedDocument,
    chunk_parsed_document,
    detect_document_type,
    parse_document_bytes,
)
from pm_agent.app.utils.source_refs import document_source_refs


def test_document_type_detection_uses_slack_metadata_and_extension():
    assert detect_document_type("spec.bin", "application/pdf", None) == "pdf"
    assert detect_document_type("spec.docx", None, None) == "docx"
    assert detect_document_type("notes.md", "text/plain", None) == "md"
    assert detect_document_type("notes", None, "text") == "txt"
    assert detect_document_type("sheet.xlsx", None, "xlsx") is None


def test_pdf_parser_preserves_pages_and_rejects_image_only_pdf():
    pdf = fitz.open()
    first = pdf.new_page()
    first.insert_text((72, 72), "First page requirement")
    second = pdf.new_page()
    second.insert_text((72, 72), "Second page clarification")
    parsed = parse_document_bytes(pdf.tobytes(), "pdf")
    pdf.close()

    assert parsed.page_count == 2
    assert parsed.pages[0][0] == 1
    assert "First page requirement" in parsed.pages[0][1]
    assert "Second page clarification" in parsed.text

    blank = fitz.open()
    blank.new_page()
    with pytest.raises(DocumentParseError, match="OCR"):
        parse_document_bytes(blank.tobytes(), "pdf")
    blank.close()


def test_docx_parser_extracts_paragraphs_and_tables():
    docx = DocxDocument()
    docx.add_paragraph("Portal requirements")
    table = docx.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Feature"
    table.cell(0, 1).text = "Invoice download"
    stream = BytesIO()
    docx.save(stream)

    parsed = parse_document_bytes(stream.getvalue(), "docx")

    assert parsed.page_count is None
    assert "Portal requirements" in parsed.text
    assert "Feature | Invoice download" in parsed.text


def test_txt_and_markdown_decode_with_replacement():
    txt = parse_document_bytes(b"hello\xffworld", "txt")
    markdown = parse_document_bytes(b"# Heading\n- item", "md")

    assert "hello�world" == txt.text
    assert markdown.text == "# Heading\n- item"


def test_chunking_is_bounded_overlapping_and_page_preserving():
    page_one = " ".join(f"alpha{i}" for i in range(700))
    page_two = " ".join(f"beta{i}" for i in range(700))
    chunks = chunk_parsed_document(
        ParsedDocument(
            text=page_one + "\n\n" + page_two,
            pages=[(1, page_one), (2, page_two)],
            page_count=2,
        )
    )

    assert len(chunks) > 2
    assert [chunk["chunk_index"] for chunk in chunks] == list(range(len(chunks)))
    assert all(len(chunk["text"]) <= CHUNK_MAX_CHARS for chunk in chunks)
    assert {chunk["page_number"] for chunk in chunks} == {1, 2}
    page_one_chunks = [chunk["text"] for chunk in chunks if chunk["page_number"] == 1]
    assert page_one_chunks[0][-80:] in page_one_chunks[1]


def test_document_context_uses_analysis_without_extracted_text():
    document = SimpleNamespace(
        id=uuid.uuid4(),
        slack_file_id="F1",
        filename="requirements.pdf",
        title=None,
        download_status="completed",
        download_error=None,
        parse_status="parsed",
        parse_error=None,
        analysis_status="completed",
        page_count=4,
        analysis_summary="Customer portal and billing requirements.",
        analysis_items={
            "requirements": [{"text": "Users download invoices"}],
            "clarifications": [{"question": "Can cancelled users download invoices?"}],
        },
        extracted_text="SECRET FULL DOCUMENT TEXT",
    )

    context = format_document_context(document, 8)

    assert "requirements.pdf" in context
    assert "Users download invoices" in context
    assert "SECRET FULL DOCUMENT TEXT" not in context


def test_gather_caps_combined_document_context_without_affecting_source_refs():
    context = gather_info_service._compact_document_contexts(["A" * 8000, "B" * 8000])

    assert len(context) <= gather_info_service.MAX_DOCUMENT_CONTEXT_PER_MESSAGE_CHARS
    assert "Additional document context omitted" in context


def test_document_source_refs_and_gather_ref_merge_are_stable():
    document = SimpleNamespace(id="d1", slack_file_id="F1")
    refs = document_source_refs(document)

    merged = gather_info_service._gather_item_source_refs(
        ["m1"],
        {"m1": {"type": "slack_message", "message_id": "m1"}},
        {"m1": refs},
    )

    assert merged == [
        {"type": "slack_message", "message_id": "m1"},
        {"type": "document", "id": "d1"},
        {"type": "slack_file", "id": "F1"},
    ]


def test_chunk_analysis_overrides_llm_source_coordinates():
    class DummyLLM:
        async def generate_json(self, **kwargs):
            return {
                "summary": "summary",
                "requirements": [
                    {
                        "text": "Audit plan changes",
                        "source": {"page": 999, "chunk_index": 999},
                        "confidence": "high",
                    }
                ],
            }

    output = asyncio.run(
        document_chunk_analysis_chain(
            text="Plan changes must be audit logged.",
            chunk_index=7,
            page_number=4,
            channel_kind="client",
            llm_client=DummyLLM(),
        )
    )

    assert output.requirements[0].source.page == 4
    assert output.requirements[0].source.chunk_index == 7


def test_analysis_merge_keeps_highest_confidence_and_unicode_items():
    low = DocumentChunkAnalysisOutput(
        requirements=[
            DocumentRequirement(
                text="صارف انوائس ڈاؤن لوڈ کر سکتے ہیں",
                confidence="low",
                source={"page": 1, "chunk_index": 1},
            )
        ]
    )
    high = DocumentChunkAnalysisOutput(
        requirements=[
            DocumentRequirement(
                text="صارف انوائس ڈاؤن لوڈ کر سکتے ہیں",
                confidence="high",
                source={"page": 2, "chunk_index": 2},
            )
        ]
    )

    merged = document_service.merge_document_analysis([low, high])

    assert len(merged["requirements"]) == 1
    assert merged["requirements"][0]["confidence"] == "high"
    assert merged["requirements"][0]["source"]["chunk_index"] == 2


def test_multi_file_registration_commits_metadata_before_scheduling(monkeypatch):
    events = []
    counter = {"value": 0}

    class FakeSession:
        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, tb):
            return False

        async def commit(self):
            events.append("commit")

    async def upsert(db, **kwargs):
        counter["value"] += 1
        return (
            SimpleNamespace(
                id=uuid.UUID(int=counter["value"]),
                download_status="pending",
                parse_status="pending",
            ),
            True,
        )

    def schedule(document_id, client=None):
        assert "commit" in events
        events.append(f"schedule:{document_id}")
        return True

    monkeypatch.setattr(document_service, "get_session_factory", lambda: lambda: FakeSession())
    monkeypatch.setattr(document_service, "upsert_document_metadata", upsert)
    monkeypatch.setattr(document_service, "schedule_document_processing", schedule)

    ids = asyncio.run(
        document_service.register_message_documents(
            project_id=uuid.uuid4(),
            channel_id="C1",
            source_message_id=uuid.uuid4(),
            files=[{"id": "F1", "name": "one.pdf"}, {"id": "F2", "name": "two.docx"}],
        )
    )

    assert len(ids) == 2
    assert events[0] == "commit"
    assert len([event for event in events if event.startswith("schedule:")]) == 2


def test_standalone_file_shared_resolves_channel_through_files_info(monkeypatch):
    project_id = uuid.uuid4()
    captured = {}

    async def find_project_channel(channel_ids):
        if "C-INTERNAL" in channel_ids:
            return project_id, "C-INTERNAL"
        return None

    async def file_info(client, file_id):
        assert file_id == "F-STANDALONE"
        return {
            "id": file_id,
            "name": "requirements.pdf",
            "mimetype": "application/pdf",
            "url_private_download": "https://files.slack.test/F-STANDALONE",
            "shares": {"private": {"C-INTERNAL": [{}]}},
        }

    async def register(**kwargs):
        captured.update(kwargs)
        return [uuid.uuid4()]

    monkeypatch.setattr(document_service, "_find_file_project_channel", find_project_channel)
    monkeypatch.setattr(document_service, "get_file_info", file_info)
    monkeypatch.setattr(document_service, "_register_documents", register)

    result = asyncio.run(
        document_service.register_file_shared_event(
            {"type": "file_shared", "file_id": "F-STANDALONE", "user_id": "U1"},
            client=object(),
        )
    )

    assert len(result) == 1
    assert captured["project_id"] == project_id
    assert captured["channel_id"] == "C-INTERNAL"
    assert captured["source_message_id"] is None
    assert captured["files"][0]["name"] == "requirements.pdf"


def test_document_task_registry_suppresses_duplicate_local_work(monkeypatch):
    document_id = uuid.uuid4()
    document_service._active_document_tasks.clear()

    async def run_test():
        release = asyncio.Event()

        async def process(document_id, client=None):
            await release.wait()

        monkeypatch.setattr(document_service, "process_document", process)
        assert document_service.schedule_document_processing(document_id) is True
        assert document_service.schedule_document_processing(document_id) is False
        release.set()
        await document_service._active_document_tasks[document_id]
        await asyncio.sleep(0)

    asyncio.run(run_test())
    assert document_id not in document_service._active_document_tasks


@pytest.mark.parametrize(
    ("document_values", "expected"),
    [
        (
            {"filename": "sheet.xlsx", "mimetype": None, "filetype": "xlsx", "size_bytes": 20},
            ("skipped", "unsupported", "skipped"),
        ),
        (
            {
                "filename": "large.pdf",
                "mimetype": "application/pdf",
                "filetype": "pdf",
                "size_bytes": document_service.MAX_DOCUMENT_BYTES + 1,
            },
            ("failed", "failed", "skipped"),
        ),
    ],
)
def test_pre_download_status_transitions_are_independent(monkeypatch, document_values, expected):
    document = SimpleNamespace(
        id=uuid.uuid4(),
        parse_status="pending",
        download_status="pending",
        analysis_status="pending",
        filename=document_values["filename"],
        mimetype=document_values["mimetype"],
        filetype=document_values["filetype"],
        size_bytes=document_values["size_bytes"],
        url_private_download="https://files.slack.test/file",
        url_private=None,
    )

    class FakeSession:
        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, tb):
            return False

        async def commit(self):
            return None

    async def get_document(db, document_id):
        return document

    async def update_document(db, target, **values):
        for key, value in values.items():
            setattr(target, key, value)
        return target

    monkeypatch.setattr(document_service, "get_session_factory", lambda: lambda: FakeSession())
    monkeypatch.setattr(document_service, "get_document", get_document)
    monkeypatch.setattr(document_service, "update_document", update_document)

    asyncio.run(document_service._process_document(document.id, object()))

    assert (document.download_status, document.parse_status, document.analysis_status) == expected
    if document.download_status == "failed":
        assert "25 MiB" in document.download_error
    else:
        assert "Unsupported" in document.parse_error


def test_download_uses_bearer_auth_and_computes_sha(monkeypatch):
    captured = {}
    payload = b"document bytes"

    class FakeResponse:
        headers = {"content-length": str(len(payload))}

        def raise_for_status(self):
            return None

        async def aiter_bytes(self):
            yield payload[:5]
            yield payload[5:]

    class FakeStream:
        async def __aenter__(self):
            return FakeResponse()

        async def __aexit__(self, exc_type, exc, tb):
            return False

    class FakeHttp:
        def __init__(self, **kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, tb):
            return False

        def stream(self, method, url, headers):
            captured.update({"method": method, "url": url, "headers": headers})
            return FakeStream()

    from pm_agent.app import config

    monkeypatch.setattr(document_service.httpx, "AsyncClient", FakeHttp)
    monkeypatch.setattr(config, "get_settings", lambda: SimpleNamespace(slack_bot_token="xoxb-test"))

    data, digest = asyncio.run(document_service._download_file("https://files.slack.test/F1"))

    assert data == payload
    assert digest == hashlib.sha256(payload).hexdigest()
    assert captured["headers"]["Authorization"] == "Bearer xoxb-test"


def test_download_rejects_slack_html_response(monkeypatch):
    payload = b"<!DOCTYPE html><html><body>Sign in to Slack</body></html>"

    class FakeResponse:
        headers = {"content-length": str(len(payload)), "content-type": "text/html"}

        def raise_for_status(self):
            return None

        async def aiter_bytes(self):
            yield payload

    class FakeStream:
        async def __aenter__(self):
            return FakeResponse()

        async def __aexit__(self, exc_type, exc, tb):
            return False

    class FakeHttp:
        def __init__(self, **kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, tb):
            return False

        def stream(self, method, url, headers):
            return FakeStream()

    from pm_agent.app import config

    monkeypatch.setattr(document_service.httpx, "AsyncClient", FakeHttp)
    monkeypatch.setattr(config, "get_settings", lambda: SimpleNamespace(slack_bot_token="xoxb-test"))

    with pytest.raises(document_service.DocumentDownloadError, match="files:read"):
        asyncio.run(document_service._download_file("https://files.slack.test/F1"))


def test_html_download_failure_marks_download_failed_not_parse_only(monkeypatch):
    document = SimpleNamespace(
        id=uuid.uuid4(),
        project_id=uuid.uuid4(),
        slack_channel_id="C",
        slack_file_id="F",
        parse_status="pending",
        parse_error=None,
        download_status="pending",
        download_error=None,
        analysis_status="pending",
        filename="spec.pdf",
        mimetype="application/pdf",
        filetype="pdf",
        size_bytes=1000,
        url_private_download="https://files.slack.test/F",
        url_private=None,
    )

    async def fake_get_document(db, document_id):
        return document

    async def fake_update_document(db, target, **values):
        for key, value in values.items():
            setattr(target, key, value)
        return target

    async def html_download(url):
        raise document_service.DocumentDownloadError(
            "Slack file bytes were not downloaded; received an HTML page instead."
        )

    class FakeSession:
        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, tb):
            return False

        async def commit(self):
            return None

    monkeypatch.setattr(document_service, "get_session_factory", lambda: lambda: FakeSession())
    monkeypatch.setattr(document_service, "get_document", fake_get_document)
    monkeypatch.setattr(document_service, "update_document", fake_update_document)
    monkeypatch.setattr(document_service, "_download_file", html_download)

    asyncio.run(document_service._process_document(document.id, object()))

    assert document.download_status == "failed"
    assert document.parse_status == "failed"
    assert document.analysis_status == "skipped"
    assert "HTML page" in document.download_error


def test_embedding_failure_leaves_failed_chunk_null(monkeypatch):
    updated = []

    class FakeSession:
        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, tb):
            return False

        async def commit(self):
            return None

    class DummyLLM:
        async def embed(self, text):
            if text == "bad":
                raise RuntimeError("embedding unavailable")
            return [0.1] * 1536

    async def update_embedding(db, chunk_id, embedding):
        updated.append(chunk_id)

    monkeypatch.setattr(document_service, "get_session_factory", lambda: lambda: FakeSession())
    monkeypatch.setattr(document_service, "update_chunk_embedding", update_embedding)

    chunks = [
        {"id": uuid.uuid4(), "text": "good"},
        {"id": uuid.uuid4(), "text": "bad"},
    ]
    asyncio.run(document_service._embed_chunks(chunks, DummyLLM()))

    assert updated == [chunks[0]["id"]]


def test_embedding_concurrency_is_capped_at_three(monkeypatch):
    active = 0
    peak = 0

    class FakeSession:
        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, tb):
            return False

        async def commit(self):
            return None

    class DummyLLM:
        async def embed(self, text):
            nonlocal active, peak
            active += 1
            peak = max(peak, active)
            await asyncio.sleep(0.001)
            active -= 1
            return [0.1] * 1536

    async def update_embedding(db, chunk_id, embedding):
        return None

    monkeypatch.setattr(document_service, "get_session_factory", lambda: lambda: FakeSession())
    monkeypatch.setattr(document_service, "update_chunk_embedding", update_embedding)
    chunks = [{"id": uuid.uuid4(), "text": str(index)} for index in range(9)]

    asyncio.run(document_service._embed_chunks(chunks, DummyLLM()))

    assert peak == document_service.DOCUMENT_CONCURRENCY


def test_full_document_analysis_processes_every_chunk(monkeypatch):
    document = SimpleNamespace(analysis_status="pending", analysis_summary=None, analysis_items={})
    seen = []

    class FakeSession:
        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, tb):
            return False

        async def commit(self):
            return None

    async def analyze(**kwargs):
        seen.append(kwargs["chunk_index"])
        return DocumentChunkAnalysisOutput(
            summary=f"chunk {kwargs['chunk_index']}",
            requirements=[
                DocumentRequirement(
                    text=f"Requirement {kwargs['chunk_index']}",
                    source={"page": kwargs["page_number"], "chunk_index": kwargs["chunk_index"]},
                )
            ],
        )

    async def summarize(outputs, items, llm_client):
        return "Merged summary"

    async def get_document(db, document_id):
        return document

    async def update_document(db, target, **values):
        for key, value in values.items():
            setattr(target, key, value)
        return target

    monkeypatch.setattr(document_service, "document_chunk_analysis_chain", analyze)
    monkeypatch.setattr(document_service, "_summarize_document", summarize)
    monkeypatch.setattr(document_service, "get_session_factory", lambda: lambda: FakeSession())
    monkeypatch.setattr(document_service, "get_document", get_document)
    monkeypatch.setattr(document_service, "update_document", update_document)

    chunks = [
        {"id": uuid.uuid4(), "chunk_index": index, "page_number": index + 1, "text": f"text {index}"}
        for index in range(4)
    ]
    asyncio.run(document_service._analyze_document(uuid.uuid4(), chunks, "client", object()))

    assert sorted(seen) == [0, 1, 2, 3]
    assert document.analysis_status == "completed"
    assert document.analysis_summary == "Merged summary"
    assert len(document.analysis_items["requirements"]) == 4


def test_analysis_failure_preserves_parsed_document_state(monkeypatch):
    document = SimpleNamespace(
        parse_status="parsed",
        analysis_status="pending",
        extracted_text="Preserved parsed text",
    )

    class FakeSession:
        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, tb):
            return False

        async def commit(self):
            return None

    async def analyze(**kwargs):
        raise RuntimeError("analysis unavailable")

    async def get_document(db, document_id):
        return document

    async def update_document(db, target, **values):
        for key, value in values.items():
            setattr(target, key, value)
        return target

    monkeypatch.setattr(document_service, "document_chunk_analysis_chain", analyze)
    monkeypatch.setattr(document_service, "get_session_factory", lambda: lambda: FakeSession())
    monkeypatch.setattr(document_service, "get_document", get_document)
    monkeypatch.setattr(document_service, "update_document", update_document)

    asyncio.run(
        document_service._analyze_document(
            uuid.uuid4(),
            [{"id": uuid.uuid4(), "chunk_index": 0, "page_number": 1, "text": "text"}],
            "client",
            object(),
        )
    )

    assert document.analysis_status == "failed"
    assert document.parse_status == "parsed"
    assert document.extracted_text == "Preserved parsed text"


def test_document_schema_and_migration_have_required_statuses_and_indexes():
    assert {"download_status", "download_error", "parse_status", "analysis_status"}.issubset(
        Document.__table__.c.keys()
    )
    assert DocumentChunk.embedding.type.dim == 1536

    migration = Path(
        "pm_agent/app/db/migrations/versions/007_document_ingestion.py"
    ).read_text()
    assert 'revision: str = "007_documents"' in migration
    assert "uq_documents_slack_file_id" in migration
    assert "ix_documents_project_download" in migration
    assert "ix_documents_project_parse" in migration
    assert "ix_documents_project_analysis" in migration
    assert "uq_document_chunks_document_index" in migration
